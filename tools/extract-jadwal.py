#!/usr/bin/env python3
"""
Ekstraksi jadwal praktek dokter dari file HFIS BPJS RS Bhirawa Bhakti.

Sumber : "JADWAL DOKTER SESUAI HFIS BPJS.xlsx" (1 sheet per poli)
Ambil  : HANYA data publik -> nama dokter, hari, jam.
Jangan : JANGAN ambil NIK / HP / nomor SIP ke output mana pun.

Struktur kolom (via merged header):
  J:K  "Jam kerja"                    -> jadwal kerja internal
  L:M  "Praktek Poli Non Eksekutif"   -> jadwal yang tampil ke publik (dipakai website)
Sebagian sheet hanya mengisi salah satu pasangan saja -> itu dilaporkan.

Output : data/jadwal-dokter.csv + .json
"""
import csv
import json
import re
import sys
from pathlib import Path

import openpyxl

SRC = Path("/home/rega/Workspace/Logo/JADWAL DOKTER SESUAI HFIS BPJS.xlsx")
OUT_DIR = Path(__file__).resolve().parent.parent / "data"

POLI_MAP = {
    "DOKTER ANAK": ("Poli Anak", ""),
    "DOKTER BEDAH": ("Poli Bedah", ""),
    "DOKTER OBSGYN": ("Poli Kandungan (Obgyn)", ""),
    "DOKTER PENYAKIT DALAM": ("Poli Penyakit Dalam", ""),
    "JANTUNG DAN PEMBULUH DARAH": ("Poli Jantung", ""),
    "SARAF": ("Poli Saraf", ""),
    "THT-KL": ("Poli THT-KL", ""),
    "GIGI BEDAH MULUT": ("Poli Gigi Spesialis", "Bedah Mulut"),
    "GIGI ENDODONSI": ("Poli Gigi Spesialis", "Endodonsi"),
    "GIGI PEDODONTIS": ("Poli Gigi Spesialis", "Pedodonsi (Gigi Anak)"),
    "GIGI PERIODONTI": ("Poli Gigi Spesialis", "Periodonsia (Gusi)"),
    "DOKTER GIGI UMUM": ("Poli Gigi Umum", ""),
    "DOKTER UMUM": ("Poli Umum", ""),
    "ANASTESI": ("Anastesi", ""),
    "RADIOLOGI": ("Radiologi", ""),
    "ORTOPHEDIA": ("Orthopedi", ""),
    "Mata": ("Poli Mata", ""),
}

HARI_VALID = ["Senin", "Selasa", "Rabu", "Kamis", "Jumat", "Sabtu", "Minggu"]
JAM_RE = re.compile(r"(\d{1,2})[.:](\d{2})\s*[-–—]\s*(\d{1,2})[.:](\d{2})")


def norm_hari(v):
    s = re.sub(r"\s+", " ", str(v or "")).strip().upper().rstrip(".")
    if not s:
        return ""
    s = s.replace("AKHAD", "MINGGU")
    if s.count("-") == 1 and "JUMAT" in s and "SENIN" in s:
        return "Senin – Jumat"
    if s in ("SETIAP HARI", "HARI KERJA"):
        return "Senin – Jumat"
    return s.title() if s in [h.upper() for h in HARI_VALID] else ""


def norm_jam(v):
    m = JAM_RE.search(str(v or ""))
    if not m:
        return "", ""
    h1, m1, h2, m2 = m.groups()
    return f"{int(h1):02d}.{m1}", f"{int(h2):02d}.{m2}"


def clean_nama(s):
    s = re.sub(r"\s+", " ", str(s or "")).strip().rstrip(",")
    s = re.sub(r",\s*", ", ", s)
    return s


def find_columns(ws):
    """Cari baris header + pasangan kolom (hari, jam) untuk jam kerja & praktek poli."""
    for row in ws.iter_rows(max_row=8):
        idx = {}
        for c in row:
            if c.value is None:
                continue
            t = str(c.value).strip().lower()
            if "nama dan gelar" in t:
                idx["nama"] = c.column
            elif t.startswith("jam kerja"):
                idx["kerja_hari"] = c.column
            elif "praktek" in t or "non eksekutif" in t:
                idx["prak_at"] = c.column
        if "nama" in idx and "kerja_hari" in idx:
            kerja_hari = idx["kerja_hari"]
            kerja_jam = kerja_hari + 1
            prak_hari = idx.get("prak_at")
            if prak_hari == kerja_jam:
                prak_hari = None  # header salah tempat, tidak ada pasangan praktek
            return {
                "row": row[0].row,
                "nama": idx["nama"],
                "kerja": (kerja_hari, kerja_jam),
                "prak": (prak_hari, prak_hari + 1) if prak_hari else None,
            }
    return None


def cell(r, col):
    return r[col - 1] if col and 0 < col <= len(r) else None


def slots(ws, pair, start_row, nama_col):
    """Ambil slot (dokter, hari, jam) dari satu pasangan kolom."""
    out = []
    dokter = ""
    for r in ws.iter_rows(min_row=start_row):
        nv = cell(r, nama_col)
        if nv is not None:
            t = "" if nv.value is None else str(nv.value).strip()
            if t and not t.isdigit():
                dokter = clean_nama(t)
        hari = norm_hari(cell(r, pair[0]).value if cell(r, pair[0]) else None)
        jam = norm_jam(cell(r, pair[1]).value if cell(r, pair[1]) else None)
        if dokter and hari and jam[0]:
            out.append({"dokter": dokter, "hari": hari, "jam": jam[0], "jam_selesai": jam[1]})
    return out


def main():
    if not SRC.exists():
        sys.exit(f"FAIL: sumber tidak ada: {SRC}")
    wb = openpyxl.load_workbook(SRC, data_only=True)
    rows, report = [], []

    for ws in wb.worksheets:
        poli, spes = POLI_MAP.get(ws.title, (ws.title.title(), ""))
        cols = find_columns(ws)
        if not cols:
            report.append(f"⚠️  {ws.title}: header tidak standar -> dilewati")
            continue
        start = cols["row"] + 2  # lewati baris header + baris nomor kolom
        prac = slots(ws, cols["prak"], start, cols["nama"]) if cols["prak"] else []
        kerj = slots(ws, cols["kerja"], start, cols["nama"])
        dipakai, sumber = (prac, "praktek") if prac else (kerj, "jam_kerja")
        nama_dokter = sorted({s["dokter"] for s in dipakai})
        for s in dipakai:
            rows.append({
                "poli": poli, "dokter": s["dokter"], "spesialisasi": spes,
                "hari": s["hari"], "jam_mulai": s["jam"],
                "jam_selesai": s["jam_selesai"], "sumber": sumber,
            })
        flag = "" if prac else "  ⟵ ⚠️ kolom 'Praktek Poli' KOSONG, pakai 'Jam kerja'"
        report.append(f"{'✅' if prac else '⚠️ '} {ws.title}: {len(dipakai)} slot / "
                      f"{len(nama_dokter)} dokter -> {poli}{flag}")
        if not dipakai and kerj:
            report.append(f"      (ada {len(kerj)} slot jam kerja tapi gagal parse)")

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    fields = ["poli", "dokter", "spesialisasi", "hari", "jam_mulai", "jam_selesai", "sumber"]
    with (OUT_DIR / "jadwal-dokter.csv").open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=fields)
        w.writeheader()
        w.writerows(rows)
    (OUT_DIR / "jadwal-dokter.json").write_text(
        json.dumps(rows, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

    print("\n".join(report))
    print(f"\nTOTAL: {len(rows)} slot | {len({r['dokter'] for r in rows})} dokter | "
          f"{len({r['poli'] for r in rows})} poli")


if __name__ == "__main__":
    main()
